import pdfplumber
from docx import Document
for fn in ['Titanic_Survival_Prediction_Project_Report.docx']:
 d=Document(fn); print(fn,'paragraphs',len(d.paragraphs),'tables',len(d.tables))
with pdfplumber.open('Titanic_Survival_Prediction_Project_Report.pdf') as pdf:
 print('PDF pages',len(pdf.pages))
 for i in [2,3,4,5,10,12,17,18,20,21]:
  print('PAGE',i+1,((pdf.pages[i].extract_text() or '').splitlines()[:4]))
