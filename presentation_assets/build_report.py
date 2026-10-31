from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]; A=Path(__file__).resolve().parent/'report_assets'; OUT=ROOT/'Titanic_Survival_Prediction_Project_Report.docx'
N='#214F84'; B='#3975BA'; G='#5AA846'; O='#DF765E'; I='#242B33'
target=A/'target.png'; eda=A/'eda.png'; missing=A/'missing.png'; modelplot=A/'models.png'; cmplot=A/'confusion.png'; importance=A/'importance.png'; ageplot=A/'age.png'
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82);sec.header_distance=Inches(.35);sec.footer_distance=Inches(.35);sec.different_first_page_header_footer=True
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string('242B33');normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
for nm,sz,col in [('Title',28,N),('Heading 1',20,N),('Heading 2',14,G),('Heading 3',11,B)]:
 st=doc.styles[nm];st.font.name='Arial';st.font.size=Pt(sz);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(col.lstrip('#'));st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
st=doc.styles.add_style('Code',WD_STYLE_TYPE.PARAGRAPH);st.font.name='Consolas';st.font.size=Pt(8);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1
def field(run,name):
 r=run._r;b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=name;s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');t=OxmlElement('w:t');t.text='1';e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end')
 for x in (b,i,s,t,e):r.append(x)
h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=h.add_run('TITANIC SURVIVAL PREDICTION  |  PAGE ');r.font.size=Pt(8);field(h.add_run(),'PAGE')
q=sec.footer.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run('Invertis University  |  B.Tech (CSE)  |  Academic project report').font.size=Pt(8)
def p(t='',lead=None,align=None):
 q=doc.add_paragraph()
 if align is not None:q.alignment=align
 if lead and t.startswith(lead):q.add_run(lead).bold=True;q.add_run(t[len(lead):])
 else:q.add_run(t)
 return q
def hd(t,n=1):return doc.add_heading(t,n)
def br():doc.add_page_break()
def bullet(t):q=doc.add_paragraph(style='List Bullet');q.paragraph_format.space_after=Pt(3);q.add_run(t)
def number(t):q=doc.add_paragraph(style='List Number');q.paragraph_format.space_after=Pt(4);q.add_run(t)
def img(path,w=6.3,alt='Project figure'):
 q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True;x=q.add_run().add_picture(str(path),width=Inches(w));x._inline.docPr.set('descr',alt)
def cap(t):
 q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(8);r=q.add_run(t);r.italic=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('59636E')
def shade(c,col):x=OxmlElement('w:shd');x.set(qn('w:fill'),col);c._tc.get_or_add_tcPr().append(x)
def cellm(c):
 m=OxmlElement('w:tcMar')
 for k in ('top','start','bottom','end'):
  x=OxmlElement('w:'+k);x.set(qn('w:w'),'110');x.set(qn('w:type'),'dxa');m.append(x)
 c._tc.get_or_add_tcPr().append(m)
def tbl(head,rows,widths=None,fs=8.5):
 t=doc.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for j,v in enumerate(head):
  c=t.rows[0].cells[j];c.text=v;shade(c,'214F84');cellm(c)
  for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(fs)
 tr=t.rows[0]._tr.get_or_add_trPr();tr.append(OxmlElement('w:tblHeader'))
 for i,row in enumerate(rows):
  cs=t.add_row().cells
  for j,v in enumerate(row):
   cs[j].text=str(v);cellm(cs[j]);cs[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if i%2:shade(cs[j],'EFF4F8')
   for q in cs[j].paragraphs:
    q.paragraph_format.space_after=Pt(0)
    for r in q.runs:r.font.size=Pt(fs)
 bd=OxmlElement('w:tblBorders')
 for edge in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'5');x.set(qn('w:color'),'D9D9D9');bd.append(x)
 t._tbl.tblPr.append(bd);return t
def code(label,snip,why):
 q=p(label);q.paragraph_format.keep_with_next=True
 for r in q.runs:r.bold=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('214F84')
 t=doc.add_table(rows=1,cols=1);c=t.cell(0,0);shade(c,'F1F4F6');c.text='';cellm(c)
 for j,line in enumerate(snip.strip().splitlines()):
  q=c.paragraphs[0] if j==0 else c.add_paragraph();q.style='Code';q.add_run(line)
 p(why)
def link(q,text,url):
 rid=q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid);run=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),B.lstrip('#'));pr.append(co);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr);tx=OxmlElement('w:t');tx.text=text;run.append(tx);el.append(run);q._p.append(el)
# Cover inspired by reference; no sample names or unrelated topics.
q=p('INVERTIS',align=WD_ALIGN_PARAGRAPH.CENTER);q.paragraph_format.space_before=Pt(30)
for r in q.runs:r.bold=True;r.font.size=Pt(34);r.font.color.rgb=RGBColor.from_string('07569B')
q=p('UNIVERSITY  BAREILLY',align=WD_ALIGN_PARAGRAPH.CENTER)
for r in q.runs:r.bold=True;r.font.size=Pt(16);r.font.color.rgb=RGBColor.from_string('C32737')
q=p('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING',align=WD_ALIGN_PARAGRAPH.CENTER);q.paragraph_format.space_before=Pt(36)
for r in q.runs:r.bold=True;r.font.size=Pt(15);r.font.color.rgb=RGBColor.from_string('5B6067')
q=p('[Confirm official department wording]',align=WD_ALIGN_PARAGRAPH.CENTER)
for r in q.runs:r.italic=True;r.font.size=Pt(9)
for a,b in [('COURSE NAME','B.Tech (CSE)'),('SECTION','[Insert section]'),('PROGRAM CODE','IIOT-4 [Confirm applicability]'),('PROGRAM NAME','Machine Learning'),('PROJECT NAME','Titanic Survival Prediction'),('MENTOR NAME','[Insert mentor name]')]:
 q=doc.add_paragraph();q.paragraph_format.left_indent=Inches(.55);q.paragraph_format.space_before=Pt(16 if a=='COURSE NAME' else 7);r=q.add_run(a+': ');r.font.size=Pt(12);r.font.color.rgb=RGBColor.from_string('3975BA');r=q.add_run(b);r.bold=True;r.font.size=Pt(12);r.font.color.rgb=RGBColor.from_string('555B61')
q=p('A supervised machine learning study of historical passenger records',align=WD_ALIGN_PARAGRAPH.CENTER);q.paragraph_format.space_before=Pt(24)
for r in q.runs:r.italic=True;r.font.size=Pt(11)
br()
hd('Project and Team Details',1);p('Complete the verified course details. Add or remove team entries as needed; no names or student IDs have been assumed.')
tbl(['Item','Details'],[('University','Invertis University, Bareilly'),('Department','Department of Computer Science and Engineering [confirm wording]'),('Course','B.Tech (CSE)'),('Section','[Insert section]'),('Program code','IIOT-4 [confirm applicability]'),('Program name','Machine Learning'),('Project','Titanic Survival Prediction'),('Mentor','[Insert mentor name]'),('Student / team name(s)','[Insert name(s)]'),('Student ID(s)','[Insert corresponding ID(s)]'),('Submission date','[Insert date]')],[2.05,4.45])
p('Confirm the official department name, program code, course section, team, mentor, and dataset citation before submission.')
hd('Abstract',2)
p('This project builds an end-to-end educational classifier using 891 Titanic passenger records. The work checks the supplied CSV, studies missing values and survival patterns, creates a family-size feature, and compares a baseline with Logistic Regression, Decision Tree, and Random Forest models. Preprocessing and classification are combined in a scikit-learn pipeline to reduce leakage during cross-validation. A stratified 80/20 split reserves 179 records for validation. Random Forest was selected by the notebook’s training cross-validation rule and achieved 79.9% accuracy and 0.833 ROC-AUC on the held-out records. The pipeline is then saved with joblib and loaded by a FastAPI website that displays an estimated probability. These estimates describe patterns in this dataset; they are not causal conclusions or guarantees about individuals.')
q=p('Keywords: supervised learning, binary classification, exploratory data analysis, Random Forest, scikit-learn, FastAPI.');q.runs[0].italic=True
br()
hd('Index',1)
for label,n in [('System Architecture Diagram','4'),('1. Introduction','5'),('2. Problem Statement','5'),('3. Objectives of the Project','5'),('4. Methodology','6'),('5. Literature Review','11'),('6. Explanation of the Project (Working)','11'),('7. Workflow','13'),('8. Program','14'),('9. Results and Discussions','16'),('10. Snippets of UI','19'),('11. Future Scope','21'),('12. Conclusion','22'),('References','22'),('Appendix A. Reproducibility Notes','22')]:
 q=doc.add_paragraph();q.paragraph_format.left_indent=Inches(.3);q.paragraph_format.space_after=Pt(7);r=q.add_run(label+'  '+'.'*max(3,64-len(label))+'  '+n);r.font.size=Pt(10.5)
br()

hd('System Architecture Diagram',1);p('This diagram summarizes the actual browser, API, saved model, notebook, and CSV components.')
img(A/'architecture.png',6.4,'Browser, FastAPI, saved pipeline, prediction response, CSV and notebook');cap('Figure 1. System architecture of the Titanic Passenger Outcome Lab.')
p('The browser serves Predict, Explore Data, and About. The form submits six fields to POST /api/predict. FastAPI validates the request, calls the saved scikit-learn pipeline, and returns an estimated probability and outcome label. GET /api/stats serves summaries from Titanic.csv. The notebook trains and saves the pipeline.')
hd('1. Introduction',1)
p('The Titanic passenger dataset contains selected passenger attributes and a recorded survival outcome. This project studies how those attributes can be used in supervised machine learning. The supplied CSV contains 891 rows. The target survived is binary: 0 means did not survive and 1 means survived.')
p('A binary classifier learns patterns from examples with known outcomes and estimates the outcome for an input profile. This project teaches data inspection, exploratory analysis, feature preparation, model evaluation, model saving, and web integration.')
p('The file contains 342 survivors (38.4%) and 549 non-survivors (61.6%). Always predicting the larger class would reach 61.6% accuracy on the whole file. The notebook therefore includes a baseline and reports metrics beyond accuracy.')
p('This is an educational analysis of one historical dataset. It does not prove causes of survival or determine what would have happened to a particular person.')
hd('2. Problem Statement',1)
p('The problem is to estimate a passenger’s recorded outcome from attributes in the CSV and collected by the website. The model uses passenger class, sex, age, family size, fare, and port of embarkation.')
p('The result is a statistical estimate from historical patterns. This project does not reconstruct a person’s history, infer unrecorded circumstances, or explain why a passenger survived. The CSV has no name, ticket number, cabin number, or passenger ID for identity lookup.')
p('The scope is a local demonstration: a notebook analyzes the data and evaluates classifiers; a complete preprocessing-plus-model pipeline is saved to joblib; FastAPI accepts example profiles through a browser interface.')
hd('3. Objectives of the Project',1);p('The project objectives are to:')
for x in ['Inspect CSV structure, types, target labels, duplicates, and missing values.','Explore survival counts and group rates using charts.','Create family_size and remove the row index and unsuitable or redundant fields.','Prevent target leakage by excluding alive, which duplicates survived.','Build preprocessing for imputation, numeric scaling, and categorical encoding.','Compare a majority-class baseline, Logistic Regression, Decision Tree, and Random Forest using stratified cross-validation.','Evaluate the selected model on a separate stratified validation set and report metrics and a confusion matrix.','Refit the selected pipeline on available rows for the demo and save with joblib.','Connect the pipeline to FastAPI and a three-page website.']:bullet(x)
hd('4. Methodology',1);hd('4.1 Dataset and source information',2)
p('Titanic.csv has 891 rows and 16 columns before cleaning. Unnamed: 0 is a row index from 0 to 890, not a passenger attribute; the notebook removes it, leaving 15 fields. These include survived, pclass, sex, age, sibsp, parch, fare, embarked, class, who, adult_male, deck, embark_town, alive, and alone.')
p('The supplied files do not identify the dataset publisher, download location, release, or license. Confirm and cite its source and reuse terms before submission; this report does not guess the provenance.')
hd('4.2 Data validation and duplicates',2)
p('The notebook checks that required fields exist and that survived contains only 0 and 1. Common fields such as Name, Ticket, Cabin, and PassengerId are absent. After removing the row index, 107 rows are identical across all remaining values. Without a passenger identifier, it is not possible to determine whether they are copies or distinct people with the same recorded values, so the notebook retains them.')
hd('Data quality checks performed',3)
tbl(['Check','Why it matters','Project decision'],[('Target values','A classifier needs a consistent label.','Verify survived is binary: 0 or 1.'),('Row index','An index can create accidental patterns.','Drop Unnamed: 0 before analysis.'),('Exact duplicate rows','Duplicates can affect counts and splits.','Report 107 identical rows; retain because identity cannot be verified.'),('Missing fields','Many missing values can distort model fit.','Measure before choosing imputation or exclusion.'),('Redundant fields','Repeated information can leak or overweight a signal.','Exclude aliases and derived columns from model inputs.')],[1.35,2.45,2.9],7.8)
hd('4.3 Missing values',2)
p('Missingness is concentrated in deck (688 rows, 77.2%), age (177, 19.9%), embarked (2, 0.2%), and embark_town (2, 0.2%). Deck is excluded because it is mostly missing and not collected by the form. Age and embarked are imputed inside the model pipeline. Redundant embark_town is excluded in favor of embarked.')
img(missing,5.0,'Missing-value percentage in deck, age, embarked, and embark_town');cap('Figure 2. Missing-value percentages in notable fields (n = 891).')
p('Deck has the most missing data and was excluded from the model inputs because the website cannot collect it reliably. Age is retained because it is available for most records and is imputed inside the pipeline. Only two rows lack embarked, so the categorical pipeline fills these using the most frequent training category.')
hd('4.4 Exploratory Data Analysis',2)
p('The notebook explores outcome balance; survival by sex and class; age and fare distributions; embarkation port; family fields; and missingness. Count charts show group sizes, while rates compare the fraction surviving. These are descriptive patterns, not causal evidence.')
hd('How to read the EDA',3)
p('The target chart shows a class imbalance: non-survivors are the larger group, so accuracy must be compared with a baseline. The sex chart shows a substantially higher observed survival rate among women in this sample. The class chart shows lower rates in second and third class than in first class. These group sizes differ, so each percentage should be read together with its n label. Age results use only rows with recorded age, and the 61+ group is especially small. No single chart controls for other passenger characteristics.')
img(target,4.1,'Target counts and percentages');cap('Figure 3. Target distribution: 342 survived (38.4%) and 549 did not (61.6%).')
p('The target is imbalanced but not extremely rare: a majority-class rule predicts “did not survive” for everyone. It provides a simple reference, not a useful survivor detector. This motivates reporting precision, recall, F1, and ROC-AUC alongside accuracy.')
img(eda,4.7,'Survival rates by sex and ticket class with sample sizes');cap('Figure 4. Observed survival rates by sex and ticket class; sample sizes are labelled.')
p('In the supplied records, the women’s observed rate is 74.2% (233/314), compared with 18.9% (109/577) for men. By ticket class, rates are 63.0% in first, 47.3% in second, and 24.2% in third class. The differences are associations in this dataset and may overlap with other attributes.')
img(ageplot,4.3,'Notebook age-band survival rates with sample sizes');cap('Figure 5. Age-band rates; the 61+ group has only 22 passengers.')
p('The 0–12 age group has an observed rate of 58.0% (40/69). The 61+ estimate is based on just 22 people, so it is less stable. Missing-age records are omitted from these age-group rates, which is why the denominators sum to fewer than 891.')
hd('4.5 Feature engineering and selection',2)
p('The notebook defines family_size = sibsp + parch + 1, where +1 counts the passenger. It combines the two family-count fields into the website feature. is_alone and family_group are created for exploration but not used in the model because they derive from family_size.')
tbl(['Column(s)','Decision','Reason'],[('survived','Target','Binary outcome to estimate.'),('alive','Exclude: leakage','Matches survived in all 891 rows and reveals the answer.'),('Unnamed: 0','Exclude','Sequential row index, not a passenger trait.'),('class, embark_town','Exclude','Repeat pclass and embarked.'),('who, adult_male, alone','Exclude','Derived or redundant information.'),('deck','Exclude','77.2% missing; absent from the form.'),('sibsp, parch','Use via family_size','Combined into one engineered field.'),('pclass, sex, age, fare, embarked','Use','Fields represented by the model form.')],[1.55,1.55,3.3],8)
hd('4.6 Split and preprocessing',2)
p('A stratified 80/20 split with random_state=42 creates 712 training and 179 validation rows. Five-fold StratifiedKFold cross-validation and tuning use training rows only. The validation rows are held aside for final comparison.')
hd('Why stratification and a pipeline are used',3)
p('Stratification preserves approximately the same survivor/non-survivor proportions in both partitions. A fixed random_state makes the split reproducible. Imputers, scaling, and category encoding are fitted inside each training fold; fitting these steps on the full dataset first would let information from validation folds influence the model and could make scores optimistic. One-hot encoding represents categories as indicator columns and ignores categories not encountered during fitting.')
tbl(['Input type','Pipeline transformation','Purpose'],[('Numeric: age, family_size, fare','Median imputation, then StandardScaler','Fill missing values and standardize numeric scale.'),('Categorical: pclass, sex, embarked','Most-frequent imputation, then one-hot encoding','Fill missing categories and create indicator columns.')],[1.55,2.3,2.55],8)
p('Pipeline and ColumnTransformer keep preprocessing attached to model fitting. Each cross-validation fold learns its imputers and scaler from its own training portion.')
hd('Metrics used for evaluation',3)
p('Accuracy is the fraction of all predictions that are correct. Precision answers: among passengers predicted to survive, how many were recorded as survivors? Recall answers: among recorded survivors, how many did the model find? F1 is the harmonic mean of precision and recall. ROC-AUC measures how well scores rank positive cases above negative cases across thresholds; 0.5 represents random ranking. The project reports all of these because the target classes are not balanced and the threshold changes the precision/recall trade-off.')
hd('4.7 Model training and selection',2)
p('The notebook compares a most-frequent DummyClassifier, Logistic Regression, a depth-limited Decision Tree, and a Random Forest. It tunes Logistic Regression C, tree depth and leaf size, and Random Forest depth and leaf size on training rows. The predeclared rule selects the highest tuned training cross-validation ROC-AUC. Random Forest is selected at 0.878. On validation, Logistic Regression scores 0.839 ROC-AUC and Random Forest 0.833, a 0.006 gap. The notebook keeps its predeclared rule and reports the difference.')
hd('4.8 Save and integrate',2)
p('The selected pipeline is refit on all 891 rows for the demo and saved as titanic_survival_pipeline.joblib. app.py loads it and checks the feature order. This all-row model has no unseen test score; reported validation metrics come from the 712/179 split. The artifact loaded successfully here using scikit-learn 1.8.0, the version pinned in requirements.txt.')
hd('5. Literature Review',1)
p('Cox’s work on regression for binary sequences describes modelling a 0/1 outcome from explanatory variables [1]. Logistic Regression is a useful binary-classification baseline, and its coefficients offer one view of associations, though encoding and feature relationships still matter.')
p('Classification trees build feature-based splits that assign records to terminal groups. The CART framework formalized tree classification and regression procedures [2]. A single Decision Tree can be described as rules, but its learned rules may vary with the data and settings.')
p('Random Forests combine many randomized tree predictors [3]. The notebook compares this ensemble with Logistic Regression and a Decision Tree. Training cross-validation selects Random Forest, while Logistic Regression scores 0.006 higher validation ROC-AUC. This is why both selection rules and held-out results should be reported.')
p('The implementation uses scikit-learn pipelines and ColumnTransformer to keep preprocessing attached to model fitting. Its model-persistence guide cautions against assuming cross-version loading will work; the environment should match the version used to save the artifact [4–7]. These sources describe methods and software, not universal model rankings.')
hd('6. Explanation of the Project (Working)',1);hd('6.1 From input to API response',2)
for x in ['The visitor opens Predict and enters class, sex, age, family size, fare, and embarkation port.','The browser sends those fields as JSON to POST /api/predict.','FastAPI validates the fields and orders them to match the saved model.','The pipeline applies preprocessing and the Random Forest classifier.','The API returns an estimated probability, higher/lower-likelihood label, model name, and cautionary note.']:number(x)
hd('6.2 Actual website inputs',2)
tbl(['API field','Meaning','Accepted values'],[('pclass','Ticket class','1, 2, or 3'),('sex','Recorded sex','male or female'),('age','Age in years','0 to 100'),('family_size','Passenger plus family aboard','1 to 20'),('fare','Ticket fare','0 to 600'),('embarked','Port code','S, C, or Q')],[1.2,2.6,2.6],8)
p('The form requires a valid value for all six fields. The pipeline imputers handle missing values during training; the API schema still expects submitted values.')
hd('Input validation and response interpretation',3)
p('The API schema bounds numeric values and restricts category fields to accepted choices so malformed profiles are rejected before inference. The response contains a percentage estimate, a label based on the application threshold, the model name, and a cautionary note. A probability is not a promise: model calibration was not separately assessed, and the historical dataset does not represent all passengers or circumstances.')
hd('6.3 Project component roles',2)
tbl(['Component','Role'],[('Titanic.csv','Notebook analysis and API summaries.'),('Titanic_Survival_Analysis.ipynb','Checks, EDA, features, comparison, validation, saving.'),('titanic_survival_pipeline.joblib','Saved preprocessing-plus-model pipeline.'),('app.py','FastAPI pages, validation, summaries, prediction.'),('static HTML/CSS/JavaScript','Predict, Explore Data, and About pages.'),('requirements.txt','Dependencies; scikit-learn pinned to 1.8.0.')],[2.65,3.75],8.5)
br();hd('7. Workflow',1);p('The notebook creates and evaluates the pipeline; the website serves estimates from the saved artifact.')
img(A/'workflow.png',5.6,'Workflow from CSV through training and web prediction');cap('Figure 6. Workflow from CSV inspection to a website estimate.')
p('The steps separate development from evaluation: first the 179-row validation set is held aside, then candidate pipelines are tuned using five folds of the 712 training rows. Reported scores come from the held-out set. After those scores are recorded, the selected pipeline is refit on all rows to power the local website demo.')
br();hd('8. Program',1);p('These excerpts come from the supplied notebook and app.py; they show selected operations, not every source file.')
code('Listing 1. Family-size feature (notebook).','''df["family_size"] = df["sibsp"] + df["parch"] + 1
ENGINEERED_FEATURES = [
  "pclass", "sex", "age", "family_size", "fare", "embarked"
]''','Family size counts the passenger and matches the website input.')
code('Listing 2. Preprocessing (notebook).','''numeric = Pipeline([
  ("impute", SimpleImputer(strategy="median")),
  ("scale", StandardScaler()),
])
categorical = Pipeline([
  ("impute", SimpleImputer(strategy="most_frequent")),
  ("onehot", OneHotEncoder(handle_unknown="ignore")),
])''','The pipeline learns preprocessing from the training portion.')
hd('8. Program (continued)',1)
code('Listing 3. Split and train-only cross-validation (notebook).','''train_df, valid_df = train_test_split(
  df, test_size=0.20, stratify=df["survived"], random_state=42
)
cv = StratifiedKFold(n_splits=5, shuffle=True, random_state=42)
search = GridSearchCV(pipe, param_grid, scoring="roc_auc", cv=cv)
search.fit(X_train, y_train)''','The validation data are kept separate from train-only tuning.')
code('Listing 4. Prediction API (app.py).','''@app.post("/api/predict")
def predict(passenger: Passenger):
  row = pd.DataFrame([passenger.model_dump()])[FEATURES]
  probability = float(model.predict_proba(row)[0, 1])
  return {
    "survival_probability": round(probability * 100, 1),
    "outcome": "Higher likelihood" if probability >= 0.5
               else "Lower likelihood",
    "model": type(model.named_steps["model"]).__name__,
  }''','The endpoint validates a profile, invokes the loaded model, and returns an estimated probability.')
code('Listing 5. Save final pipeline (notebook).','''final_model = clone(selected_model).fit(
  df[WEBSITE_FEATURES], df["survived"]
)
joblib.dump(final_model, PROJECT_DIR / "titanic_survival_pipeline.joblib")''','The demo pipeline is refit on all rows; its output is not a held-out score.')

br();hd('9. Results and Discussions',1);hd('9.1 Dataset and descriptive results',2)
tbl(['Measure','Observed result'],[('Rows','891'),('Did not survive','549 (61.6%)'),('Survived','342 (38.4%)'),('Age missing','177 (19.9%)'),('Deck missing','688 (77.2%)'),('Embarked missing','2 (0.2%)'),('Identical rows after index removal','107; retained because copy status is unknown.')],[2.7,3.7],8.5)
p('Observed survival was 74.2% for women (233/314) and 18.9% for men (109/577). By class it was 63.0% (136/216) in first, 47.3% (87/184) in second, and 24.2% (119/491) in third. Notebook age-band rates were 58.0% for ages 0–12 (40/69), 42.9% for 13–18 (30/70), 38.3% for 19–35 (137/358), 40.0% for 36–60 (78/195), and 22.7% for 61+ (5/22). Age rates omit 177 rows with missing age. These are descriptive associations.')
hd('9.2 Candidate model comparison',2)
p('Random Forest is selected by its tuned training cross-validation ROC-AUC (0.878). Validation ROC-AUC is 0.839 for Logistic Regression and 0.833 for Random Forest. The notebook keeps its training-CV selection rule and discloses the small validation difference.')
tbl(['Model','Accuracy','Precision','Recall','F1','ROC-AUC'],[('Logistic Regression','0.799','0.780','0.667','0.719','0.839'),('Random Forest (selected)','0.799','0.811','0.623','0.705','0.833'),('Decision Tree','0.765','0.765','0.565','0.650','0.785'),('Dummy baseline','0.615','0.000','0.000','0.000','0.500')],[1.8,.95,.85,.7,.7,.8],7.6)
img(modelplot,5.8,'Validation ROC-AUC by model');cap('Figure 7. Validation ROC-AUC; 0.5 is the random-ranking reference.')
hd('9.3 Selected model metrics',2)
tbl(['Metric','Random Forest','Meaning'],[('Accuracy','79.9% (143/179)','Share of all records classified correctly.'),('Precision','81.1% (43/53)','Share of predicted survivors who survived.'),('Recall','62.3% (43/69)','Share of actual survivors found.'),('F1','0.705','Balance between precision and recall.'),('ROC-AUC','0.833','Ranking across classification thresholds.')],[1.1,1.4,3.5],8)
p('The majority-class baseline scored 61.5% accuracy, 0.000 F1, and 0.500 ROC-AUC. Accuracy alone would hide that it never identifies survivors.')
hd('9.4 Error patterns and feature importance',2)
p('The confusion matrix and permutation plot provide complementary diagnostics: the first counts errors at the chosen threshold, while the second shows how shuffling each input changes validation ROC-AUC.')
img(A/'evaluation_panels.png',5.95,'Confusion matrix and permutation importance diagnostics')
cap('Figure 8. Confusion matrix (left) and Figure 9. Permutation importance (right), computed on validation data.')
p('The matrix contains 100 correct non-survivor predictions, 10 false alarms, 26 missed survivors, and 43 correctly found survivors. Missed survivors explain why recall (62.3%) is lower than precision (81.1%). A different classification threshold can raise recall while also raising false positives; threshold choice should follow a clearly stated use case and must not be tuned on a final test set.')
p('Permutation importance measures how validation ROC-AUC changes when a feature is shuffled. Sex has the largest mean decrease (0.2247), followed by class (0.0346) and age (0.0334); family_size has the smallest value (0.0099). Correlated fields can share or mask importance, and results vary across samples and random shuffles. This is a model diagnostic, not a causal explanation.')
hd('9.5 Validation model and demo model',2)
p('The validation model uses 712 training rows and is scored on 179 unseen rows. The saved model is then refit on all 891 rows for the website. Its displayed probability is a model estimate, not an independent test score.')
br();hd('10. Snippets of UI',1);p('These are genuine browser captures of the supplied local FastAPI pages. The ship hero uses generated illustrative artwork, not an archival photograph.')
for i,(title,fn,caption,desc) in enumerate([('10.1 Predict page','predict.png','Figure 10. Predict page.','Visitors enter six fields and receive a class label and estimated probability.'),('10.2 Explore Data page','explore.png','Figure 11. Explore Data page.','The page presents summary values, group charts, controls, and a field guide.'),('10.3 About page','about.png','Figure 12. About page.','The page describes the pipeline, input fields, interpretation, and limitations.')]):
 if i:br()
 hd(title,2);img(ROOT/'presentation_assets'/fn,6.1,caption);cap(caption);p(desc)
hd('11. Future Scope',1);p('The following suggestions are future work, not implemented features:')
for x in ['Confirm and document the original dataset source, version, license, and citation.','Repeat evaluation across multiple stratified splits or repeated cross-validation.','Calibrate probabilities and explain the effect of the 0.5 threshold on precision and recall.','Improve keyboard accessibility, labels, and mobile responsiveness.','Add careful per-example explanations without implying causation.','Deploy with environment controls, monitoring, and documented model updates.']:bullet(x)
hd('12. Conclusion',1)
p('This project completes an educational machine learning workflow using the supplied passenger records. It validates the CSV, explores survival patterns and missingness, creates family_size, prevents target leakage, compares classifiers, evaluates a selected Random Forest, saves a complete pipeline, and serves estimates through FastAPI.')
p('The selected model reached 79.9% accuracy, 0.705 F1, and 0.833 ROC-AUC on 179 validation rows. It found 43 of 69 survivors and missed 26. Logistic Regression had slightly higher validation ROC-AUC (0.839) than Random Forest (0.833), while Random Forest remained selected under the notebook’s training-CV rule.')
p('These results describe one split of one historical dataset. Missing values, limited attributes, class imbalance, small sample size, and unverified provenance limit interpretation. The saved demo is refit on all rows; outputs are estimates, not guarantees or explanations of historical outcomes.')
hd('References',1);p('Research and software documentation used to describe the methods. Project values come from the supplied CSV, notebook, application, and artifact.')
refs=[('[1] Cox, D. R. (1958). “The Regression Analysis of Binary Sequences.” JRSS: Series B, 20(2), 215–232.','https://doi.org/10.1111/j.2517-6161.1958.tb00292.x'),('[2] Breiman, L., Friedman, J. H., Olshen, R. A., & Stone, C. J. (1984). Classification and Regression Trees. Wadsworth. ISBN 978-0-412-04841-8.','https://books.google.com/books?id=JwQx-WOmSyQC'),('[3] Breiman, L. (2001). “Random Forests.” Machine Learning, 45, 5–32.','https://doi.org/10.1023/A:1010933404324'),('[4] scikit-learn developers. Logistic Regression, User Guide.','https://scikit-learn.org/stable/modules/linear_model.html#logistic-regression'),('[5] scikit-learn developers. Decision Trees, User Guide.','https://scikit-learn.org/stable/modules/tree.html'),('[6] scikit-learn developers. Model Persistence.','https://scikit-learn.org/stable/model_persistence.html'),('[7] scikit-learn developers. Pipelines and Composite Estimators.','https://scikit-learn.org/stable/modules/compose.html'),('[8] Dataset source, version, access date, and license: [To be confirmed by student].','')]
for label,url in refs:
 q=doc.add_paragraph();q.paragraph_format.left_indent=Inches(.25);q.paragraph_format.first_line_indent=Inches(-.25);q.paragraph_format.space_after=Pt(7);q.add_run(label+' ')
 if url:
  rid=q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid);run=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=url;run.append(tx);el.append(run);q._p.append(el)
hd('Appendix A. Reproducibility Notes',1)
p('Start the application from the project directory using README.md. It runs locally at http://127.0.0.1:8000. Pages are /predict, /explore, and /about; API routes include GET /api/stats and POST /api/predict.')
p('The saved artifact loaded successfully here using scikit-learn 1.8.0. It is a Pipeline with preprocess and model steps and expects pclass, sex, age, family_size, fare, embarked. requirements.txt pins scikit-learn==1.8.0.')
p('Charts were generated from Titanic.csv. Metrics and permutation importance were read from executed notebook outputs. UI images are genuine local-page captures. The dataset source remains unverified.')
p('Files inspected: Titanic.csv; Titanic_Survival_Analysis.ipynb; app.py; titanic_survival_pipeline.joblib; requirements.txt; static/index.html; static/explore.html; static/about.html; static/styles.css; static/pages.css; static/app.js; static/explore.js; static/images/titanic-dawn.png.')
doc.core_properties.title='Titanic Survival Prediction Project Report';doc.core_properties.subject='Invertis University machine learning project';doc.core_properties.author='[Student Name]'
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd);doc.save(OUT);print(OUT)



